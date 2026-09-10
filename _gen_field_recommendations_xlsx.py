"""Generate Echo AI Book field recommendation Excel from the Echo field map."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SRC = Path(r"c:\Users\Administrator\Downloads\02_air-import-job-fields_2_fixed_EN.xlsx")
OUT = Path(r"D:\work\testenv\airfreight-western-ui\Echo-AI-Book-Field-Recommendations.xlsx")
OUT2 = Path(r"c:\Users\Administrator\Downloads\Echo-AI-Book-Field-Recommendations.xlsx")

HAVE = {
    "customerId",
    "shipperId",
    "consigneeId",
    "notifyId",
    "carrierId",
    "bookingAgent",
    "designationAgent",
    "loadingPortCode",
    "dischargingPortCode",
    "etd",
    "eta",
    "quantityActual",
    "weightActual",
    "cubageActual",
    "chargeWeight",
    "cargoEnglishName",
    "packing",
    "cargoType",
    "hbl",
    "mbl",
    "cargoSource",
    "incoTerm",
    "freightTerm",
    "paymentTermHBL",
    "paymentTermHbl",
    "opOffice",
    "opDepartment",
    "sales",
    "hsCode",
}

PRIO_A = {
    "operateType",
    "voyageFlight",
    "vessel",
    "customsBroker",
    "paymentTermMBL",
    "op",
    "remarks",
    "satisfiedRequests",
    "customs",
    "hsCode",
}

PRIO_B = {
    "loadingAgentId",
    "groundAgentId",
    "exChangeAgent",
    "cargoCode",
    "mark",
    "measurements",
    "tradeTerm",
    "deliveryAddress",
    "salesOffice",
    "salesDepartment",
    "quotationNo",
    "declarantBroker",
    "customerContact",
    "contactTelephone",
    "contactEmail",
    "ata",
    "atd",
    "bookingDate",
    "deliverDate",
    "masterJobNo",
    "cargoSalesId",
}

STAFF_LATER = {
    "csr",
    "csrOffice",
    "csrDepartment",
    "dc",
    "dcOffice",
    "dcDepartment",
    "planner",
    "plannerOffice",
    "plannerDepartment",
    "overseaOp",
    "overseaOpDepartment",
    "siteOP",
    "loadingPortOp",
    "destinationPortOp",
}

ORDER = {
    "KEEP — use now": 0,
    "ADD — Priority A": 1,
    "ADD — Priority B": 2,
    "LATER": 3,
    "SKIP — Priority C": 4,
    "SKIP": 5,
}

FILLS = {
    "KEEP — use now": PatternFill("solid", fgColor="C6EFCE"),
    "ADD — Priority A": PatternFill("solid", fgColor="FCE4D6"),
    "ADD — Priority B": PatternFill("solid", fgColor="FFF2CC"),
    "LATER": PatternFill("solid", fgColor="DDEBF7"),
    "SKIP": PatternFill("solid", fgColor="E7E6E6"),
    "SKIP — Priority C": PatternFill("solid", fgColor="F8CBAD"),
}


def bucket_display(disp: str) -> str:
    s = (disp or "").strip()
    sl = s.lower()
    if not s or sl.startswith("n/a") or s.upper() == "N/A":
        return "N/A / not suggested"
    if "form display" in sl:
        return "Form Display (suggested)"
    if "system" in sl:
        return "System"
    if "not displayed" in sl:
        return "Not displayed"
    return "Other"


def advice(field: str, cat: str, disp_bucket: str, in_ui: bool) -> tuple[str, str]:
    if disp_bucket != "Form Display (suggested)":
        if disp_bucket == "System":
            return "SKIP", "System-owned / not operator-entered on create"
        if disp_bucket.startswith("N/A"):
            return "SKIP", "Echo marked N/A — not for OS create UI"
        return "SKIP", "Not Form Display"

    if "Other LOB" in cat or "Warehouse" in cat:
        return "SKIP", "Out of AU AI create scope (warehouse / other LOB)"
    if "UDF" in cat or "userDefined" in field:
        return "SKIP — Priority C", "Tenant UDF — only if customer requires"
    if "Customized" in cat:
        return "SKIP", "OJ/customized tenant fields"

    if field in HAVE or field == "cargoEnglishName":
        return "KEEP — use now", "Already on Book / Clearance; correct for trial create"

    if field in PRIO_A:
        return "ADD — Priority A", "High value for AU AI first-trial create; add next"

    if field in PRIO_B:
        return "ADD — Priority B", "Add when create deepens"

    if cat.startswith("Staff & Dept"):
        if field.startswith("userDefinedRole"):
            return "SKIP — Priority C", "Custom posts — tenant-specific"
        if field in STAFF_LATER:
            return "LATER", "Full staff matrix — default later, not first-screen create"
        return "ADD — Priority B", "Useful when create deepens"

    if cat.startswith("Customs") or "Tax" in cat or cat == "Charge & Finance":
        return "LATER", "Clearance / finance / post-create — not Book essentials"

    if cat in ("Service Term", "Sales Channel", "Statics", "Flight"):
        return "LATER", "Optional ops detail — second wave"

    if cat.startswith("JOB Basic - Identification"):
        return "LATER", "Assigned/system or cross-company — not typed on create"

    if in_ui:
        return "KEEP — use now", "Already on Book / Clearance"

    if cat.startswith("Party - Agent") or cat.startswith("Contact") or cat.startswith(
        "Party - Customer"
    ):
        return "ADD — Priority B", "Form Display but secondary for first trial"

    if cat in (
        "Cargo",
        "Cargo & QWC",
        "Freight & Trade Terms",
        "Timeline",
        "JOB Basic - MBL/HBL",
        "Place",
        "Remarks & Notes",
    ):
        return "ADD — Priority B", "Form Display but secondary for first trial"

    return "LATER", "Form Display in full H5; defer for trial create"


def apply_rec_fill(cell, recommendation: str) -> None:
    if recommendation in FILLS:
        cell.fill = FILLS[recommendation]
        return
    for key, fill in FILLS.items():
        if recommendation.startswith(key.split("—")[0].strip()):
            cell.fill = fill
            return


def main() -> None:
    src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws_in = src["01_Air Import Job"]

    rows: list[dict] = []
    for i, row in enumerate(ws_in.iter_rows(max_col=11, values_only=True), 1):
        if i == 1:
            continue
        if i > 900:
            break
        if not row or row[1] is None:
            continue
        cat = str(row[0] or "")
        field = str(row[1] or "")
        typ = str(row[2] or "")
        meaning = str(row[3] or "").replace("\n", " ")
        legacy_disp = str(row[5] or "")
        os_field = str(row[7] or "")
        display = str(row[8] or "")
        label = str(row[10] or row[9] or "")
        db = bucket_display(display)
        in_ui = field in HAVE or field in ("paymentTermHbl", "cargoEnglishName")
        rec, note = advice(field, cat, db, in_ui)
        rows.append(
            {
                "category": cat,
                "field": field,
                "type": typ,
                "label": label,
                "echo_display": display,
                "display_bucket": db,
                "legacy_displayed": legacy_disp,
                "os_field": os_field,
                "in_book_ui_now": "YES" if in_ui else "NO",
                "recommendation": rec,
                "rationale": note,
                "meaning": meaning[:300],
            }
        )

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )

    # Sheet 1: All fields
    ws1 = wb.active
    ws1.title = "All fields recommendation"
    headers = [
        "Category",
        "Field",
        "Type",
        "Label",
        "Echo Display or not",
        "Display bucket",
        "Legacy Displayed",
        "OS Field",
        "In Book UI now",
        "Recommendation",
        "Rationale",
        "Meaning (short)",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws1.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    rows_sorted = sorted(
        rows,
        key=lambda r: (
            0 if r["display_bucket"].startswith("Form") else 1,
            ORDER.get(r["recommendation"], 9),
            r["category"],
            r["field"],
        ),
    )

    for r_i, r in enumerate(rows_sorted, 2):
        vals = [
            r["category"],
            r["field"],
            r["type"],
            r["label"],
            r["echo_display"],
            r["display_bucket"],
            r["legacy_displayed"],
            r["os_field"],
            r["in_book_ui_now"],
            r["recommendation"],
            r["rationale"],
            r["meaning"],
        ]
        for c, v in enumerate(vals, 1):
            cell = ws1.cell(r_i, c, v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = thin
        apply_rec_fill(ws1.cell(r_i, 10), r["recommendation"])

    widths = [28, 28, 12, 28, 40, 22, 14, 36, 14, 20, 48, 40]
    for i, w in enumerate(widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.auto_filter.ref = f"A1:L{len(rows_sorted) + 1}"
    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 30

    # Sheet 2: Summary
    ws2 = wb.create_sheet("Summary")
    c_rec = Counter(r["recommendation"] for r in rows)
    c_disp = Counter(r["display_bucket"] for r in rows)
    c_form = Counter(
        r["recommendation"] for r in rows if r["display_bucket"].startswith("Form")
    )

    ws2["A1"] = "CargoWare OS — Echo Form Display field advice"
    ws2["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws2["A2"] = (
        "Source: 02_air-import-job-fields_2_fixed_EN.xlsx · sheet 01_Air Import Job"
    )
    ws2["A3"] = (
        "Scope: AU Air Import Book / create trial vs full legacy H5 Form Display"
    )

    ws2["A5"] = "Echo Display bucket counts (all mapped fields)"
    ws2["A5"].font = Font(bold=True)
    ws2["A6"] = "Display bucket"
    ws2["B6"] = "Count"
    ws2["A6"].fill = header_fill
    ws2["B6"].fill = header_fill
    ws2["A6"].font = header_font
    ws2["B6"].font = header_font
    for i, (k, v) in enumerate(sorted(c_disp.items()), 7):
        ws2.cell(i, 1, k)
        ws2.cell(i, 2, v)

    ws2["A13"] = "Recommendation counts (ALL fields)"
    ws2["A13"].font = Font(bold=True)
    ws2["A14"] = "Recommendation"
    ws2["B14"] = "Count"
    ws2["A14"].fill = header_fill
    ws2["B14"].fill = header_fill
    ws2["A14"].font = header_font
    ws2["B14"].font = header_font
    for i, (k, v) in enumerate(sorted(c_rec.items(), key=lambda x: ORDER.get(x[0], 9)), 15):
        ws2.cell(i, 1, k)
        ws2.cell(i, 2, v)
        apply_rec_fill(ws2.cell(i, 1), k)

    ws2["A23"] = (
        "Recommendation counts (Form Display only — the real suggested UI list)"
    )
    ws2["A23"].font = Font(bold=True)
    ws2["A24"] = "Recommendation"
    ws2["B24"] = "Count"
    ws2["A24"].fill = header_fill
    ws2["B24"].fill = header_fill
    ws2["A24"].font = header_font
    ws2["B24"].font = header_font
    for i, (k, v) in enumerate(
        sorted(c_form.items(), key=lambda x: ORDER.get(x[0], 9)), 25
    ):
        ws2.cell(i, 1, k)
        ws2.cell(i, 2, v)
        apply_rec_fill(ws2.cell(i, 1), k)

    ws2["A34"] = "How to read Recommendation"
    ws2["A34"].font = Font(bold=True)
    guide = [
        ("KEEP — use now", "Already on Book/Clearance and correct for trial create."),
        (
            "ADD — Priority A",
            "Missing but high value — add next for real AU AI create.",
        ),
        (
            "ADD — Priority B",
            "Form Display; add when create deepens (agents, dims, sales office…).",
        ),
        ("LATER", "Valid on full job page; not needed on first Book screen."),
        (
            "SKIP / Priority C",
            "System, N/A, warehouse, UDF, or tenant-only — do not put on create.",
        ),
    ]
    for i, (a, b) in enumerate(guide, 35):
        ws2.cell(i, 1, a)
        ws2.cell(i, 2, b)
        apply_rec_fill(ws2.cell(i, 1), a if a in FILLS else a.split("/")[0].strip())
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 72

    # Sheet 3: Priority A
    ws3 = wb.create_sheet("Priority A — add next")
    h3 = ["Field", "Label", "Category", "OS Field", "Rationale", "In Book UI now"]
    for c, h in enumerate(h3, 1):
        cell = ws3.cell(1, c, h)
        cell.fill = PatternFill("solid", fgColor="C65911")
        cell.font = header_font
    prio_a_rows = [r for r in rows if r["recommendation"] == "ADD — Priority A"]
    for r_i, r in enumerate(sorted(prio_a_rows, key=lambda x: x["field"]), 2):
        for c, v in enumerate(
            [
                r["field"],
                r["label"],
                r["category"],
                r["os_field"],
                r["rationale"],
                r["in_book_ui_now"],
            ],
            1,
        ):
            cell = ws3.cell(r_i, c, v)
            cell.fill = FILLS["ADD — Priority A"]
    for i, w in enumerate([24, 28, 28, 40, 50, 14], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # Sheet 4: KEEP
    ws4 = wb.create_sheet("KEEP — use now")
    for c, h in enumerate(
        ["Field", "Label", "Category", "OS Field", "In Book UI now"], 1
    ):
        cell = ws4.cell(1, c, h)
        cell.fill = PatternFill("solid", fgColor="548235")
        cell.font = header_font
    keep_rows = [r for r in rows if r["recommendation"] == "KEEP — use now"]
    for r_i, r in enumerate(
        sorted(keep_rows, key=lambda x: (x["category"], x["field"])), 2
    ):
        for c, v in enumerate(
            [
                r["field"],
                r["label"],
                r["category"],
                r["os_field"],
                r["in_book_ui_now"],
            ],
            1,
        ):
            cell = ws4.cell(r_i, c, v)
            cell.fill = FILLS["KEEP — use now"]
    for i, w in enumerate([24, 28, 28, 40, 14], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    # Sheet 5: Verdict
    ws5 = wb.create_sheet("Verdict")
    ws5["A1"] = "Verdict for AU Air Import Book trial"
    ws5["A1"].font = Font(bold=True, size=14, color="1F4E79")
    verdict = [
        "",
        "Echo Form Display list = ~125 fields (full legacy job page).",
        "Do NOT put all Form Display fields on the first Book screen.",
        "",
        "~26 KEEP fields already on Book/Clearance = correct core set.",
        "~8–10 Priority A fields are the right next adds for a real create trial.",
        "Priority B = deepen create later (extra agents, dims, sales office…).",
        "LATER / SKIP = customs-tax, warehouse, UDF, full staff matrix, system IDs.",
        "",
        'Filter tip: on sheet "All fields recommendation", filter column Recommendation.',
    ]
    for i, line in enumerate(verdict, 2):
        ws5.cell(i, 1, line)
    ws5.column_dimensions["A"].width = 100

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    wb.save(OUT2)

    print("saved", OUT)
    print("saved", OUT2)
    print("total rows", len(rows))
    print("KEEP", sum(1 for r in rows if r["recommendation"] == "KEEP — use now"))
    print("ADD A", sum(1 for r in rows if r["recommendation"] == "ADD — Priority A"))
    print("ADD B", sum(1 for r in rows if r["recommendation"] == "ADD — Priority B"))
    print("LATER", sum(1 for r in rows if r["recommendation"] == "LATER"))
    print("SKIP", sum(1 for r in rows if r["recommendation"].startswith("SKIP")))


if __name__ == "__main__":
    main()
